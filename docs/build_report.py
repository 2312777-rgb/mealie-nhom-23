from pathlib import Path
from datetime import datetime
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output'
ASSETS = ROOT / 'docs' / 'assets'
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'
SHA = '879b133b5f98dd91419948f363d2d6d58793c3a0'


def diagram() -> Path:
    im = Image.new('RGB', (1500, 900), 'white')
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(FONT, 27)
    b = ImageFont.truetype(BOLD, 29)
    def box(x: int, y: int, w: int, h: int, title: str, body: str) -> None:
        d.rounded_rectangle((x, y, x+w, y+h), 12, fill='#eef3f8', outline='#324f6b', width=3)
        d.text((x+w/2, y+28), title, font=b, fill='black', anchor='mt')
        for i, line in enumerate(body.split('\n')):
            d.text((x+w/2, y+75+i*35), line, font=f, fill='black', anchor='mt')
    def arrow(a: tuple[int, int], z: tuple[int, int]) -> None:
        d.line((a, z), fill='#324f6b', width=4)
        x,y = z
        if a[0] == x:
            d.polygon([(x,y),(x-10,y-18),(x+10,y-18)], fill='#324f6b')
        else:
            d.polygon([(x,y),(x-18,y-10),(x-18,y+10)], fill='#324f6b')
    box(35, 35, 330, 170, 'Người dùng và QA', 'Thao tác nghiệp vụ\nDuyệt ảnh và phân loại diff')
    box(455, 35, 420, 170, 'Nuxt 4 và Vue 3', 'UI và API client\nlocalhost 3000')
    box(965, 35, 490, 170, 'Playwright Chromium', '8 màn hình và 3 smoke test\nBaseline và HTML JSON')
    arrow((365,120),(455,120))
    d.line((965,155,875,155), fill='#324f6b', width=4)
    d.polygon([(875,155),(893,145),(893,165)], fill='#324f6b')
    box(455, 295, 420, 175, 'FastAPI trong Docker', 'Controller và dependency\nlocalhost 19000')
    arrow((665,205),(665,295))
    d.text((685,240), 'REST và xác thực', font=f, fill='black')
    box(35, 555, 390, 160, 'Service', 'Nghiệp vụ khi cần\nMeal plan và shopping list')
    box(510, 555, 390, 160, 'Repository và HttpRepo', 'Truy vấn SQLAlchemy\nPhạm vi group household')
    box(1010, 555, 445, 160, 'SQLite trong volume QA', 'Công thức và phân loại\nThực đơn và shopping list')
    d.line((560,470,560,510,230,510,230,555), fill='#324f6b', width=4)
    arrow((230,520),(230,555))
    arrow((705,470),(705,555))
    arrow((425,635),(510,635))
    arrow((900,635),(1010,635))
    d.text((750,800), 'CRUD đơn giản có thể đi trực tiếp từ controller qua HttpRepo tới repository', font=f, fill='black', anchor='mt')
    path=ASSETS/'architecture.png'; im.save(path)
    return path


def para(doc: Document, text: str, bold: bool=False) -> None:
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.line_spacing=1.12
    r=p.add_run(text); r.bold=bold


def heading(doc: Document, title: str, level: int=1) -> None:
    doc.add_heading(title, level=level)


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate([headers]+rows):
        cells=t.rows[0].cells if row_index==0 else t.add_row().cells
        for i, (cell, text) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5)
            p.paragraph_format.line_spacing=1.05
            r=p.add_run(str(text)); r.font.name='Arial'; r.font.size=Pt(10)
            r.bold=row_index==0
            if row_index==0: r.font.color.rgb=RGBColor(255,255,255)
            pr=cell._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd'); shading.set(qn('w:fill'), '29465B' if row_index==0 else ('F0F4F7' if row_index%2==0 else 'FFFFFF')); pr.append(shading)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def page(doc: Document, title: str) -> None:
    doc.add_page_break(); heading(doc,title)


def code(doc: Document, text: str) -> None:
    for line in text.split('\n'):
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
        r=p.add_run(line); r.font.name='Consolas'; r.font.size=Pt(9)


def preview(source: Path) -> Path:
    with Image.open(source) as image:
        # Report shows a readable top-of-page excerpt; original full-page PNG stays in evidence.
        image=image.crop((0,0,image.width,min(image.height,900)))
        target=ASSETS/(source.stem+'-preview.png')
        image.save(target)
    return target


def build(final: bool) -> Path:
    doc=Document(); sec=doc.sections[0]
    # Remove inherited Word template rules beneath the Title and Subtitle styles.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.3); sec.right_margin=Cm(2)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        s=doc.styles[name]; s.font.name='Times New Roman'; s.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(12)
    doc.styles['Title'].font.size=Pt(24)
    doc.styles['Heading 1'].font.size=Pt(16)
    doc.styles['Heading 2'].font.size=Pt(13)
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].paragraph_format.space_after=Pt(9)
    header=sec.header.paragraphs[0]; header.text='R02 K10   |   Mealie   |   Nhóm 23   |   4 thành viên'
    header.runs[0].font.size=Pt(9); header.runs[0].font.color.rgb=RGBColor(0,0,0)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('Trang '); r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    p=doc.add_paragraph('TRƯỜNG ĐẠI HỌC ĐÀ LẠT\nKHOA CÔNG NGHỆ THÔNG TIN'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold=True; p.runs[0].font.size=Pt(14)
    doc.add_paragraph('\n\n')
    p=doc.add_paragraph('BÁO CÁO CUỐI KỲ' if final else 'BÁO CÁO GIỮA KỲ',style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('MÔN KIỂM THỬ PHẦN MỀM'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=True
    p=doc.add_paragraph('Kiểm thử hồi quy giao diện hệ thống Mealie bằng Playwright',style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size=Pt(19)
    doc.add_paragraph('\n')
    for text in ['Nhóm 23','Tổ hợp đề tài R02 và K10','Giảng viên Nguyễn Thế Lâm','Quy mô 4 sinh viên','Ngày cập nhật 20 tháng 10 năm 2026','Phiên bản hệ thống Mealie 3.28.0']:
        p=doc.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('\n')
    para(doc,'Hồ sơ kỹ thuật được tổng hợp từ mã nguồn và kết quả kiểm chứng trong môi trường QA cục bộ. Thông tin định danh nhóm, họ tên, MSSV và xác nhận đóng góp cần được nhóm điền trước khi nộp.')

    page(doc,'Mục lục')
    chapters=[
        'Tổng quan và phạm vi đề tài','Kiến trúc hệ thống','Cấu trúc mã nguồn và dependency',
        'Luồng dữ liệu của ba nghiệp vụ','Môi trường và khả năng tái lập','Hướng dẫn dựng và chạy',
        'Cơ sở và tiêu chí áp dụng K10','Mô hình và kịch bản kiểm thử','Kết quả kiểm chứng và bằng chứng',
    ]
    if final: chapters.append('Phân tích sai lệch và nguyên nhân')
    chapters += ['Quản trị nhóm và kế hoạch','Đối chiếu yêu cầu và kết luận']
    table(doc,['Mục','Nội dung','Trang'],[[str(i+1),title,str(i+3)] for i,title in enumerate(chapters)],[1.6,13.5,1.6])
    para(doc,'Phụ lục phía sau phần kết luận cung cấp ảnh baseline của các màn hình và minh chứng ba luồng nghiệp vụ. Mã nguồn, ảnh gốc và report tương tác được bàn giao cùng bộ hồ sơ.')

    page(doc,'1 Tổng quan và phạm vi đề tài')
    para(doc,'Nhóm áp dụng Visual Regression Testing cho Mealie để phát hiện các thay đổi giao diện ngoài ý muốn. Sản phẩm gồm môi trường chạy có dữ liệu riêng, bộ kiểm thử Playwright cho tám màn hình, ba bài smoke test nghiệp vụ, ảnh baseline và báo cáo kết quả. Phạm vi được giới hạn để có thể dựng lại và giải thích rõ khi bảo vệ.')
    para(doc,'R02 là hệ thống quản lý công thức và thực đơn Mealie. K10 yêu cầu so sánh screenshot với baseline, bảo vệ ít nhất tám màn hình và phân biệt thay đổi có chủ đích với regression. Các yêu cầu chung A đến H trong tài liệu môn học được đối chiếu ở cuối báo cáo.')
    table(doc,['Actor','Nghiệp vụ liên quan'],[
        ['Người dùng','Đăng nhập, tìm và xem công thức, xem lịch thực đơn, quản lý danh sách mua sắm'],
        ['Người tổ chức dữ liệu','Quản lý category và food theo quyền canOrganize'],
        ['Quản trị viên','Quản lý tài khoản và cấu hình hệ thống'],
        ['Nhóm QA','Khóa phiên bản, tạo seed, duyệt baseline, chạy test và phân tích diff'],
    ],[4,12.7])
    heading(doc,'Thông tin nhóm',2)
    table(doc,['Thành viên','Họ tên và MSSV','Phạm vi phân công đề xuất'],[
        ['TV1','Chưa cung cấp','Phân tích kiến trúc và tổng hợp báo cáo'],
        ['TV2','Chưa cung cấp','Docker và dữ liệu QA và hướng dẫn dựng lại'],
        ['TV3','Chưa cung cấp','Playwright VR01 đến VR04 và tự động hóa'],
        ['TV4','Chưa cung cấp','VR05 đến VR08 và phân tích diff và demo'],
    ],[2,5,9.7])
    para(doc,'Bảng phân công của Nhóm 23 ghi nhận vai trò và phạm vi bàn giao của bốn thành viên. Link commit/task và đánh giá chéo cần được nhóm bổ sung bằng dữ liệu thực tế nếu giảng viên yêu cầu.')

    page(doc,'2 Kiến trúc hệ thống')
    para(doc,'Môi trường QA sử dụng frontend Nuxt trên host ở cổng 3000 và backend FastAPI trong Docker ở cổng host 19000. SQLAlchemy lưu dữ liệu SQLite trong volume QA riêng. Trong production, upstream có thể phục vụ frontend SPA đã build qua FastAPI; cấu hình đó không phải môi trường được dùng trong lần kiểm chứng này.')
    doc.add_picture(str(diagram()),width=Cm(16.7))
    para(doc,'Hình 1 Kiến trúc container và component của môi trường QA',True)
    table(doc,['Module','Vai trò và quan hệ với kiểm thử'],[
        ['Pages và composables','Hiển thị và quản lý trạng thái UI; trực tiếp tạo nội dung ảnh chụp'],
        ['API client và xác thực','Gọi REST API; giữ phiên đăng nhập và thông tin người dùng'],
        ['Controller và schema','Nhận HTTP, kiểm tra input, dùng dependencies để xác định user và context'],
        ['Service và repository','Xử lý nghiệp vụ khi cần; CRUD đơn giản có thể qua HttpRepo tới repository'],
        ['SQLAlchemy và DB','Lưu dữ liệu; seed ổn định giúp so sánh ảnh có ý nghĩa'],
    ],[4.5,12.2])
    para(doc,'Dữ liệu được giới hạn theo group và household tùy loại. Công thức có thể được dùng ở phạm vi group; meal plan và shopping list gắn household. Bộ test dùng group và household QA riêng để giảm nhiễu từ dữ liệu cá nhân.')

    page(doc,'3 Cấu trúc mã nguồn và dependency')
    table(doc,['Đường dẫn','Nội dung cần hiểu'],[
        ['frontend/app/pages','Page Vue và định tuyến tám màn hình'],
        ['frontend/app/components','Component Domain và global và Layout'],
        ['frontend/app/composables','Logic dùng chung và xác thực và meal plan'],
        ['frontend/app/lib/api','Client API và các type sinh tự động'],
        ['mealie/routes','Controller theo recipe và household và group và auth'],
        ['mealie/services và mealie/repos','Nghiệp vụ và truy cập dữ liệu'],
        ['mealie/schema và mealie/db','Pydantic schema và SQLAlchemy model'],
        ['mealie/alembic','Migration cơ sở dữ liệu'],
        ['mealie-nhom-23','Harness và cấu hình và seed và tài liệu và bằng chứng'],
    ],[7,9.7])
    heading(doc,'Các dependency chính',2)
    para(doc,'Backend dùng Python 3.14, FastAPI, Pydantic, SQLAlchemy, Alembic và Uvicorn. Dependency được lấy từ uv.lock bằng uv sync --frozen. Frontend dùng Nuxt 4, Vue 3, TypeScript và Vuetify; phiên bản giải quyết được khóa bởi pnpm-lock.yaml. Harness dùng Playwright Test 1.64.0 và TypeScript, cài bằng npm ci theo package-lock.json.')
    para(doc,'Không chỉnh thủ công frontend/app/lib/api/types. Không thay đổi schema hoặc migration của upstream trong phạm vi này. Công việc bổ sung nằm trong harness nhằm quan sát hệ thống và cung cấp bằng chứng kiểm thử.')
    para(doc,'Source khảo sát: mealie/routes/_base/base_controllers.py, mealie/routes/_base/mixins.py, mealie/repos/all_repositories.py, frontend/app/lib/api/base/base-clients.ts và frontend/app/composables/use-mealie-auth.ts.')

    page(doc,'4 Luồng dữ liệu của ba nghiệp vụ')
    heading(doc,'Luồng tìm và xem công thức',2)
    para(doc,'Người dùng mở recipe finder và chọn bộ lọc. Page hoặc composable gọi API client; route recipe hoặc explore nhận query tùy ngữ cảnh hiển thị. Controller chuyển pagination/filter cho repository, repository lấy dữ liệu trong phạm vi hợp lệ và trả schema. Vue render card và điều khiển tìm kiếm. VR01 bảo vệ bố cục trang finder, còn FLOW01 tạo công thức qua API rồi mở trang chi tiết để xác nhận dữ liệu xuất hiện.')
    code(doc,'Người dùng → Vue page → API client → Recipe hoặc Explore controller\n→ Repository → SQLAlchemy → SQLite → Schema → Vue render')
    heading(doc,'Luồng lập thực đơn',2)
    para(doc,'Frontend gửi POST /api/households/mealplans với date, entryType và recipeId hoặc title. GroupMealplanController kiểm tra CreatePlanEntry, gắn group/user từ phiên đăng nhập rồi dùng HttpRepo và repository để lưu. Controller phát sự kiện mealplan_entry_created. Frontend đọc entry theo khoảng ngày và cập nhật lịch. FLOW02 tạo một bữa trưa QA rồi xác nhận bữa đó xuất hiện trên planner; dữ liệu tạm được xóa sau test.')
    code(doc,'Planner → POST mealplans → CreatePlanEntry → GroupMealplanController\n→ HttpRepo → RepositoryMeals → SQLite → Entry và sự kiện → Planner')
    heading(doc,'Luồng danh sách mua sắm',2)
    para(doc,'Người dùng tạo danh sách hoặc thêm nguyên liệu từ công thức. Controller shopping list dùng HttpRepo cho CRUD danh sách; ShoppingListService xử lý nghiệp vụ liên quan recipe/item. Repository lưu dữ liệu theo household và hệ thống phát sự kiện cập nhật. FLOW03 tạo danh sách qua API rồi đọc trang danh sách trong UI; VR05 bảo vệ bố cục trang tổng hợp.')
    code(doc,'Shopping UI → REST lists hoặc items → Controller\n→ HttpRepo hoặc ShoppingListService → Repository → SQLite → UI')
    para(doc,'Ba smoke test là luồng kết hợp API ghi dữ liệu và UI đọc dữ liệu. Chúng không được trình bày như các bài thao tác hoàn toàn bằng form UI, và chưa bao phủ mọi nhánh nghiệp vụ.')

    page(doc,'5 Môi trường và khả năng tái lập')
    table(doc,['Thành phần','Cấu hình'],[
        ['Mealie','3.28.0 từ commit khóa trong project-lock.json'],
        ['Backend','Linux container với Python 3.14 và uv 0.11.7'],
        ['Frontend','Nuxt dev trên host và API_URL=http://localhost:19000'],
        ['DB','SQLite trong volume mealie-r02-k10_qa-data'],
        ['UI và API','http://localhost:3000 và http://localhost:19000'],
        ['Trình duyệt','Chromium của Playwright 1.64.0 trên Windows'],
        ['Ảnh chụp','1440 × 900 và scale factor 1 và light theme'],
        ['Ngôn ngữ và thời gian','en-US và UTC và Date cố định 2026-10-12T12:00:00Z'],
    ],[4.6,12.1])
    para(doc,'Commit được pin: '+SHA+'. Tài khoản QA được bootstrap trong group R02 K10 QA và household riêng. Credential nằm trong .env; trạng thái đăng nhập nằm trong .auth và cả hai đều bị loại khỏi Git và gói nộp.')
    heading(doc,'Dữ liệu kiểm thử',2)
    para(doc,'Seed phiên bản 1 tạo 10 công thức QA Recipe 01 đến 10, 3 category, 10 food, hai danh sách QA Weekly Shopping và QA Pantry Shopping và 7 bữa tối trong tuần cố định. Category được gắn vào công thức để trang phân loại có dữ liệu. Script dùng API, kiểm tra tên hoặc khóa ngày/công thức/loại bữa trước khi tạo lại; không xóa dữ liệu hiện hữu.')
    para(doc,'Khả năng tái lập cần đồng thời khóa commit, lockfile, seed, OS, browser, font và cấu hình chụp. Baseline Windows không được coi là tương đương baseline Linux. Các dependency hệ điều hành trong image vẫn có thể thay đổi khi rebuild; lưu image ID và metadata môi trường để xác định build đã dùng.')

    page(doc,'6 Hướng dẫn dựng và chạy')
    para(doc,'Bắt đầu từ checkout Mealie đã pin, đặt harness trong thư mục mealie-nhom-23. Bật Docker Desktop với Linux containers. Các lệnh dưới đây chạy trên terminal PowerShell; xem docs/onboarding.md để có hướng dẫn đầy đủ và cách xử lý lỗi.')
    heading(doc,'Khởi động backend và frontend',2)
    code(doc,'# Từ gốc Mealie\ndocker compose -f mealie-nhom-23/docker-compose.visual.yml up -d --build --wait\n# Terminal thứ hai trong frontend\npnpm install --frozen-lockfile\n$env:API_URL="http://localhost:19000"\npnpm run dev --no-fork')
    heading(doc,'Khởi tạo tài khoản và kiểm thử',2)
    code(doc,'# Trong mealie-nhom-23\nCopy-Item .env.example .env\nnpm ci\nnpx playwright install chromium\n# Điền URL localhost và credential QA và MEALIE_ALLOW_SEED=true\nnpm run bootstrap\nnpm run seed\nnpm run typecheck\nnpm run test:smoke\nnpm run test:visual:update\n# Duyệt đủ 8 ảnh rồi mới chấp nhận baseline\nnpm run test:visual\nnpm run test:demo\nnpm run test:visual:report')
    para(doc,'Không tự cập nhật baseline khi test fail. Xác định trước lỗi do môi trường, dữ liệu, thay đổi có chủ đích hay regression. Lệnh so sánh thông thường dùng updateSnapshots=none nên không tự chấp nhận ảnh bị thiếu.')
    para(doc,'Lỗi đã gặp khi dựng native Windows: python-ldap yêu cầu Microsoft C++ Build Tools và thư viện LDAP. Cách xử lý được chọn là backend Linux container có build-essential, libldap2-dev và libsasl2-dev; giữ nguyên mã xác thực upstream.')

    page(doc,'7 Cơ sở và tiêu chí áp dụng K10')
    para(doc,'Visual Regression Testing dùng ảnh đã được duyệt làm test oracle. Với cùng trạng thái dữ liệu và môi trường, ảnh mới được so sánh với baseline để nhận biết thay đổi hình ảnh. Kỹ thuật phù hợp với màn hình ổn định và có thể kiểm soát viewport, dữ liệu, thời gian và browser.')
    heading(doc,'Quy trình kiểm thử',2)
    para(doc,'Mỗi test đăng nhập bằng tài khoản QA, mở đúng route, xác nhận nội dung đại diện đã hiển thị, đợi skeleton biến mất và font/ảnh tải xong. Sau đó test tắt animation, transition và caret để chụp toàn trang. API HTTP từ 400 trở lên, lỗi request API và lỗi JavaScript được ghi nhận để tránh lấy trang hỏng làm baseline.')
    table(doc,['Điều kiện','Quyết định'],[
        ['Pass','Đúng route và dữ liệu đại diện; không có lỗi API/JavaScript; ảnh trong ngưỡng 1% pixel khác biệt'],
        ['Fail','Sai route hoặc thiếu dữ liệu hoặc lỗi API/JavaScript hoặc ảnh vượt ngưỡng'],
        ['Baseline thiếu','Fail khi chạy so sánh; tạo ảnh bằng lệnh update riêng và duyệt trực quan'],
        ['Intentional change','Có yêu cầu và người duyệt chấp thuận; lưu lý do trước khi cập nhật baseline'],
        ['Regression','Thay đổi chưa được chấp thuận; phân tích nguyên nhân và sửa trước khi chạy lại'],
    ],[4.5,12.2])
    para(doc,'Ngưỡng maxDiffPixelRatio=0.01 là cấu hình thực nghiệm của nhóm, không phải tiêu chuẩn chung của môn học và không phải độ bao phủ mã nguồn. Sai lệch nhỏ ở chi tiết quan trọng vẫn có thể bị bỏ qua, nên cần review diff và cân nhắc assertion theo vùng khi mở rộng.')
    para(doc,'Comparator còn dùng threshold=0.2 cho mức khác biệt cảm nhận màu trên từng pixel. Trong thử nghiệm, đổi nền trắng sang xanh rất nhạt chưa bị gắn cờ; đổi sang xanh đậm tạo mismatch. Vì vậy tỷ lệ 1% không có nghĩa mọi thay đổi màu trên hơn 1% diện tích đều bị phát hiện.')
    para(doc,'Giới hạn: bộ ảnh không chứng minh đúng mọi nghiệp vụ, bảo mật, accessibility hoặc hiệu năng. Trang finder có thể kéo dài theo lượng dữ liệu; ảnh full-page chịu ảnh hưởng của lazy loading. CI và máy tạo baseline phải cùng nền tảng trước khi so sánh có ý nghĩa.')

    page(doc,'8 Mô hình và kịch bản kiểm thử')
    table(doc,['ID','Màn hình và input','Precondition và expected result'],[
        ['VR01','Recipe finder và group slug','Chưa chọn food/tool; No recipes found và bộ lọc khớp ảnh'],
        ['VR02','Recipe timeline và ngày cố định','Có QA Recipe 01; timeline và tile giữ bố cục'],
        ['VR03','Recipe categories','Category QA Breakfast gắn recipe; grid và điều khiển giữ bố cục'],
        ['VR04','Meal planner và tuần cố định','Có QA Recipe trong tuần; lịch và entry giữ bố cục'],
        ['VR05','Shopping lists','Có QA Weekly Shopping và QA Pantry Shopping; list card và navigation giữ bố cục'],
        ['VR06','User profile','Đã đăng nhập; Welcome và tóm tắt tài khoản xuất hiện'],
        ['VR07','Group foods','Có QA Food 01; bảng food và điều khiển giữ bố cục'],
        ['VR08','Meal plan settings','Có quyền truy cập; heading và setting control giữ bố cục'],
    ],[1.5,5,10.2])
    para(doc,'Input dùng chung gồm URL localhost, group slug, storageState của QA account, seed phiên bản 1 và ngày cố định. Expected result cuối cùng của từng VR là ảnh baseline cùng tên trong tests/__screenshots__/visual.spec.ts. Thư mục test-results chứa actual/diff khi có sai lệch.')
    table(doc,['ID','Nghiệp vụ bổ sung','Kết quả quan sát'],[
        ['FLOW01','API tạo công thức rồi UI mở chi tiết','Tên công thức tạm xuất hiện và dữ liệu được dọn'],
        ['FLOW02','API tạo bữa trưa rồi UI mở planner','Entry tạm xuất hiện trong tuần cố định và được dọn'],
        ['FLOW03','API tạo shopping list rồi UI mở danh sách','Tên list tạm xuất hiện và dữ liệu được dọn'],
    ],[1.8,7,7.9])

    page(doc,'9 Kết quả kiểm chứng và bằng chứng')
    result_path=ROOT/'evidence'/'verification-summary.json'
    results=json.loads(result_path.read_text(encoding='utf-8-sig')) if result_path.exists() else {}
    para(doc,'Các kết quả dưới đây lấy từ lần chạy tự động được lưu trong thư mục evidence. Kết quả pass là bằng chứng cho phạm vi đã chọn trong cấu hình này; không suy ra toàn bộ hệ thống không còn lỗi.')
    rows=results.get('rows', [['TypeScript','Đã kiểm tra thành công','tsc --noEmit'],['Danh sách test','12 test được phát hiện','8 visual và 3 smoke và 1 setup'],['Chạy ứng dụng và baseline','Chưa tổng hợp','Cần xem evidence trước khi nộp']])
    table(doc,['Hạng mục','Kết quả','Bằng chứng'],rows,[4.1,5.1,7.5])
    para(doc,results.get('conclusion','Phần kết quả ứng dụng phải được cập nhật từ lần chạy thực tế; không có số liệu suy đoán.'))
    heading(doc,'Cách đọc kết quả',2)
    para(doc,'HTML report hiển thị expected, actual và diff cho test sai lệch. JSON ghi trạng thái từng test và thời gian chạy. Seed manifest mô tả số lượng dữ liệu đã tạo. Metadata môi trường ghi commit, phiên bản runtime, image và cấu hình chụp. Ảnh của ba smoke test là bằng chứng API ghi dữ liệu và UI đọc dữ liệu.')
    para(doc,'Tỷ lệ màn hình có kịch bản là 8 trên 8 màn hình được chọn. Đây là độ bao phủ phạm vi UI đã định nghĩa, không phải statement/branch coverage của toàn bộ source. Nếu không phát hiện lỗi tự nhiên, bằng chứng demo chèn CSS có kiểm soát được dùng để đánh giá khả năng phát hiện thay đổi, và được ghi rõ là lỗi tổng hợp.')

    if final:
        page(doc,'10 Phân tích sai lệch và nguyên nhân')
        para(doc,'Hai demo chèn CSS chỉ vào trang trình duyệt để tạo thay đổi quan sát được. Chúng không sửa source Mealie và không được trình bày là lỗi phát hiện trong upstream. Sau demo, baseline gốc được giữ lại và bộ test bình thường được chạy lại để xác nhận trở về trạng thái ban đầu.')
        table(doc,['Mã','Tình huống','Phân loại và nguyên nhân'],[
            ['DEMO R01','Dịch main sang phải 48 pixel','Regression tổng hợp; CSS transform làm thay đổi vị trí nội dung và bị so sánh ảnh phát hiện'],
            ['DEMO I01','Đổi nền main sang màu xanh đậm','Intentional tổng hợp theo REQ-DEMO-01; vẫn tạo diff và cần duyệt trước khi đổi baseline'],
        ],[2.2,5.6,8.9])
        para(doc,'Luồng xử lý: tái hiện trên cùng môi trường → xem expected/actual/diff → đối chiếu yêu cầu → xác định nguyên nhân CSS/component/data → sửa lỗi hoặc xin duyệt thay đổi → chạy lại. Không dùng lệnh update-snapshots để làm mất bằng chứng regression.')
        para(doc,'Khi tạo defect thật, ghi ID, build SHA, scenario, precondition, bước tái hiện, expected/actual, mức độ, ảnh diff, nguyên nhân và trạng thái. Defect log của nhóm phân biệt rõ lỗi harness/môi trường với lỗi sản phẩm. Ngưỡng 1% có thể bỏ qua khác biệt nhỏ; đây là hạn chế cần được nêu khi bảo vệ.')
        demo=ROOT/'evidence'/'demo-regression'
        diffs=list(demo.rglob('*diff.png')) if demo.exists() else []
        if diffs:
            doc.add_picture(str(diffs[0]),width=Cm(15.5)); para(doc,'Hình 2 Diff từ demo CSS dịch vị trí nội dung',True)

    page(doc,('11' if final else '10')+' Quản trị nhóm và kế hoạch')
    para(doc,'Nhóm có bốn người. Bảng dưới là khung triển khai và bàn giao. Do chưa có danh sách tên/MSSV và lịch sử công việc do nhóm cung cấp, báo cáo không gán thành tích hoặc điểm đóng góp cho cá nhân.')
    table(doc,['Vai trò','Công việc','Tiêu chí bàn giao'],[
        ['TV1','Kiến trúc và data flow và báo cáo','Giải thích được ba nghiệp vụ và quan hệ module'],
        ['TV2','Môi trường Docker và seed và onboarding','Dựng lại từ máy sạch và seed không nhân bản'],
        ['TV3','VR01 đến VR04 và cấu hình và CI','Giải thích assertion và pixel threshold và baseline'],
        ['TV4','VR05 đến VR08 và defect và demo','Phân loại diff và trình bày root cause'],
    ],[2,7,7.7])
    heading(doc,'Các mốc theo tài liệu môn học',2)
    table(doc,['Mốc','Nội dung'],[
        ['20 tháng 10','Demo hệ thống; sơ đồ kiến trúc; onboarding; kế hoạch K10 và tiêu chí đánh giá'],
        ['Trước bảo vệ một tuần','Nộp mã nguồn và báo cáo tổng kết đã rà soát thông tin nhóm'],
        ['Ngày thi cuối kỳ','Demo automation và phân tích diff; trả lời câu hỏi; đánh giá chéo'],
    ],[4.2,12.5])
    para(doc,'Các việc nhóm phải xác nhận trước khi nộp: tên nhóm và họ tên/MSSV; lịch sử Git/board thực tế; reviewer duyệt baseline; ngày thi cuối kỳ và yêu cầu hình thức từ giảng viên. Ngày bảo vệ cụ thể chưa có trong tài liệu, không tự suy đoán.')
    para(doc,'Kịch bản demo đề xuất: giới thiệu phạm vi → mở Mealie và ba luồng nghiệp vụ → giải thích kiến trúc → chạy smoke → chạy 8 visual → chèn CSS demo → xem diff và root cause → trình bày cách khôi phục và hạn chế. CI validate được cấu hình cho GitHub hosted runner; chạy visual dùng Windows self-hosted runner cùng nền tảng baseline và cần được đăng ký trước.')

    page(doc,('12' if final else '11')+' Đối chiếu yêu cầu và kết luận')
    table(doc,['Phần','Sản phẩm và vị trí'],[
        ['A','Mục tiêu và actor và cây module và dependency trong chương 1 đến 4'],
        ['B','Sơ đồ component/container và năm module trong chương 2'],
        ['C','Môi trường và seed và hướng dẫn dựng trong chương 5 và 6; bằng chứng tại evidence'],
        ['D','Nguyên lý và phạm vi và rủi ro và pass fail trong chương 7'],
        ['E','Tám visual và ba smoke với precondition/expected trong chương 8'],
        ['F','tests và scripts và package.json và README'],
        ['G','Report và metrics và demo diff và phân tích nguyên nhân trong evidence'],
        ['H','project-lock.json và lockfile và onboarding và metadata môi trường'],
    ],[1.6,15.1])
    para(doc,'Bộ hồ sơ giới hạn phạm vi đúng theo R02 và K10: kiểm thử hồi quy giao diện của Mealie. Giá trị kiểm chứng nằm ở việc chạy lại trong cùng môi trường, sử dụng ảnh đã duyệt và giải thích được mỗi sai lệch. Bằng chứng tổng hợp không được nhầm với defect tự nhiên; kết quả chưa chạy không được ghi thành đã đạt.')
    heading(doc,'Tài liệu và mã nguồn tham khảo',2)
    para(doc,'1. Hướng dẫn thực hiện đề tài môn Kiểm thử phần mềm của giảng viên Nguyễn Thế Lâm, tài liệu PDF do nhóm cung cấp, trang 1 đến 4 và mục R02 trang 5 và K10 trang 9 đến 10.')
    para(doc,'2. Mealie source repository tại https://github.com/mealie-recipes/mealie, commit '+SHA+'. Các đường dẫn module được ghi trong docs/architecture.md.')
    para(doc,'3. Tài liệu Playwright Visual Comparisons tại https://playwright.dev/docs/test-snapshots. Cấu hình thực tế của nhóm xem playwright.config.ts và tests/visual.spec.ts.')
    para(doc,'4. Hồ sơ chạy cục bộ trong evidence và hướng dẫn tái lập docs/onboarding.md. Thời điểm tổng hợp là ngày 12 tháng 10 năm 2026.')

    shots=sorted((ROOT/'tests'/'__screenshots__').rglob('VR*.png'))
    if shots:
        limit=8 if final else 4
        for i in range(0,min(len(shots),limit),2):
            page(doc,'Phụ lục ảnh giao diện '+str(i//2+1))
            for shot in shots[i:i+2]:
                para(doc,shot.stem.replace('-',' '),True)
                shot=preview(shot)
                # Cap tall full-page screenshots; keep legibility in the original PNG evidence.
                with Image.open(shot) as img:
                    w,h=img.size
                width=min(15.2,8.0*w/h)
                doc.add_picture(str(shot),width=Cm(width))
                para(doc,'Trích phần đầu ảnh toàn trang. Xem ảnh gốc trong tests/__screenshots__/visual.spec.ts để đọc đầy đủ.')
    flows=sorted((ROOT/'evidence'/'flows').glob('FLOW*.png'))
    for i in range(0,len(flows),2):
        page(doc,'Phụ lục minh chứng nghiệp vụ '+str(i//2+1))
        for shot in flows[i:i+2]:
            para(doc,shot.stem+' API ghi dữ liệu và UI đọc dữ liệu',True)
            shot=preview(shot)
            with Image.open(shot) as img: w,h=img.size
            doc.add_picture(str(shot),width=Cm(min(15.2,8.0*w/h)))
            para(doc,'Ảnh chụp tự động khi smoke test xác nhận dữ liệu tạm xuất hiện. Dữ liệu tạm được xóa sau test.')
    path=OUT/('Bao_cao_cuoi_ky_R02_K10_Mealie.docx' if final else 'Bao_cao_giua_ky_R02_K10_Mealie.docx')
    doc.save(path)
    return path


if __name__=='__main__':
    for final in [False,True]: print(build(final))
